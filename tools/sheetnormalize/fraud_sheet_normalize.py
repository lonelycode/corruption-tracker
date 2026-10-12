#!/usr/bin/env python3
# SPDX-License-Identifier: AGPL-3.0-only
# SPDX-FileCopyrightText: 2026 Martin Buhr and contributors

"""One-off normalizer for a fraud/bribery/corruption research spreadsheet.

Reads an xlsx workbook and emits normalized JSON (by default
sources/fraud_sheet_normalized.json):
one object per data row with typed/cleaned fields plus the raw originals and
a `problems` list. Rows with problems are NOT dropped — the Go importer
(`ingester fraud import`) routes them to the human review queue instead of
auto-importing them. This script never guesses: anything ambiguous is a
problem, not a silent fix.

The output contains personal data and lives under sources/ (git-ignored).

Usage:
    python3 corruption-tracker/tools/sheetnormalize/fraud_sheet_normalize.py \
        --in sources/research.xlsx [--sheet NAME] \
        [--out sources/fraud_sheet_normalized.json]

The sheet's header row must carry the column names in COLS below. --sheet
picks the worksheet; it defaults to the workbook's active sheet.

Requires openpyxl (pip install openpyxl).
"""

import argparse
import datetime as dt
import json
import re
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl required: pip install openpyxl (or use a venv)")

COLS = {
    "year_convicted": "Year Convicted",
    "conviction_date": "Date of conviction",
    "agency": "Prosecuting agency",
    "region": "NZ Geographical area",
    "first_name": "First name",
    "last_name": "Last Name",
    "gender": "Male/Female",
    "penalty": "Penalty",
    "profession": "Profession",
    "area": "Area of Fraud",
    "amount": "$ Amount ",
    "conviction_desc": "Conviction",
    "appeal": "Appeal?",
    "company": "Company involved",
    "organisation": "Organisation",
    "victims": "Victims",
    "details": "Details",
    "link": "Link",
    "prev_convictions": "Previous/subsequent convictions",
    "aliases": "Aliases",
    "co_accused": "co-accused",
}

URL_RE = re.compile(r"^https?://", re.I)


def clean_str(v):
    if v is None:
        return ""
    s = str(v).replace("\xa0", " ").strip()
    return re.sub(r"\s+", " ", s)


def parse_year(v, problems):
    """Year Convicted: int year, or junk (URLs, prose) -> None + problem."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        y = int(v)
        if 1990 <= y <= 2030:
            return y
        problems.append(f"year_convicted out of range: {y!r}")
        return None
    s = clean_str(v)
    if re.fullmatch(r"\d{4}", s):
        y = int(s)
        if 1990 <= y <= 2030:
            return y
    problems.append(f"year_convicted not a year: {s[:60]!r}")
    return None


DATE_FORMATS = ["%d.%m.%Y", "%d.%m.%y", "%b %d %Y", "%d %b %Y", "%d/%m/%Y", "%d/%m/%y"]


def parse_date(v, problems):
    """Date of conviction -> (iso_date|None, precision|None)."""
    if v is None:
        return None, None
    if isinstance(v, (dt.datetime, dt.date)):
        d = v.date() if isinstance(v, dt.datetime) else v
        if d.year < 1990:
            # A bare year typed into a date-formatted cell: Excel stored the
            # serial day N (e.g. 2020 -> 1905-07-12). Recover the year.
            serial = (d - dt.date(1899, 12, 30)).days
            if 1990 <= serial <= 2030:
                return f"{serial}-01-01", "year"
            problems.append(f"conviction_date pre-1990 artifact: {d.isoformat()}")
            return None, None
        return d.isoformat(), "day"
    if isinstance(v, (int, float)):
        y = int(v)
        if 1990 <= y <= 2030:
            return f"{y}-01-01", "year"
        problems.append(f"conviction_date numeric but not a year: {v!r}")
        return None, None
    s = clean_str(v)
    if not s:
        return None, None
    if re.fullmatch(r"\d{4}", s):
        return f"{s}-01-01", "year"
    for fmt in DATE_FORMATS:
        try:
            d = dt.datetime.strptime(s, fmt)
            # %y two-digit years land in 1900s for >68 — clamp to 20xx.
            if d.year < 2000:
                d = d.replace(year=d.year + 100)
            return d.date().isoformat(), "day"
        except ValueError:
            continue
    # "Mar 2024" / "October 2024" style.
    for fmt in ("%b %Y", "%B %Y"):
        try:
            d = dt.datetime.strptime(s, fmt)
            return d.date().replace(day=1).isoformat(), "month"
        except ValueError:
            continue
    problems.append(f"conviction_date unparseable: {s[:60]!r}")
    return None, None


def parse_amount(v, problems):
    """$ Amount -> float NZD or None. Never guesses on ambiguity."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        a = float(v)
        return a if a > 0 else None
    s = clean_str(v)
    if not s or s.lower() in ("n/a", "na", '"', "-", "unknown"):
        return None
    t = s.replace("$", "").replace(",", "").rstrip(".")
    t = t.strip()
    # A clean number (one optional decimal point) parses; anything else —
    # prose, multi-dot typos like '1.128.233.11' — is a problem.
    if re.fullmatch(r"\d+(\.\d{1,2})?", t):
        return float(t)
    problems.append(f"amount unparseable: {s[:80]!r}")
    return None


def normalize_gender(v):
    s = clean_str(v).lower()
    return s if s in ("m", "f") else ""


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--in", dest="infile", required=True,
                    help="the research spreadsheet (.xlsx)")
    ap.add_argument("--sheet", default="",
                    help="worksheet name (default: the workbook's active sheet)")
    ap.add_argument("--out", dest="outfile", default="sources/fraud_sheet_normalized.json")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.infile, data_only=True)
    ws = wb[args.sheet] if args.sheet else wb.active
    header = [c.value for c in ws[1]]
    idx = {}
    for key, col in COLS.items():
        if col not in header:
            sys.exit(f"column {col!r} not found in sheet header")
        idx[key] = header.index(col)

    out = []
    for rownum, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        raw = {k: row[i] for k, i in idx.items()}
        if all(v is None or clean_str(v) == "" for v in raw.values()):
            continue  # fully blank row

        problems = []
        year = parse_year(raw["year_convicted"], problems)
        date_iso, date_prec = parse_date(raw["conviction_date"], problems)
        amount = parse_amount(raw["amount"], problems)

        first = clean_str(raw["first_name"])
        last = clean_str(raw["last_name"])
        full_name = (first + " " + last).strip()
        # A URL in the name/year columns marks a shifted/junk row.
        if URL_RE.match(full_name) or URL_RE.match(clean_str(raw["year_convicted"])):
            problems.append("row appears shifted (URL in name/year column)")
        if not full_name:
            problems.append("no person name")

        # Cross-check: if both present and the year disagrees with the date,
        # trust the full date but flag it.
        if year and date_iso and abs(int(date_iso[:4]) - year) > 1:
            problems.append(f"year_convicted {year} disagrees with conviction_date {date_iso}")

        link = clean_str(raw["link"])
        if link and not URL_RE.match(link):
            problems.append(f"link not a URL: {link[:60]!r}")
            link = ""

        out.append({
            "row": rownum,
            "full_name": full_name,
            "gender": normalize_gender(raw["gender"]),
            "year_convicted": year,
            "conviction_date": date_iso,
            "conviction_date_precision": date_prec,
            "amount_nzd": amount,
            "agency_raw": clean_str(raw["agency"]),
            "region_raw": clean_str(raw["region"]),
            "penalty_raw": clean_str(raw["penalty"]),
            "profession_raw": clean_str(raw["profession"]),
            "area_raw": clean_str(raw["area"]),
            "conviction_desc": clean_str(raw["conviction_desc"]),
            "appeal_raw": clean_str(raw["appeal"]),
            "company_raw": clean_str(raw["company"]),
            "organisation_raw": clean_str(raw["organisation"]),
            "victims_raw": clean_str(raw["victims"]),
            "details": clean_str(raw["details"]),
            "link": link,
            "prev_convictions_raw": clean_str(raw["prev_convictions"]),
            "aliases_raw": clean_str(raw["aliases"]),
            "co_accused_raw": clean_str(raw["co_accused"]),
            "amount_raw": clean_str(raw["amount"]),
            "year_raw": clean_str(raw["year_convicted"]),
            "date_raw": clean_str(raw["conviction_date"]),
            "problems": problems,
        })

    with open(args.outfile, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=1)

    n_problems = sum(1 for r in out if r["problems"])
    print(f"rows: {len(out)}  clean: {len(out) - n_problems}  with-problems: {n_problems}")
    from collections import Counter
    kinds = Counter(p.split(":")[0] for r in out for p in r["problems"])
    for k, n in kinds.most_common():
        print(f"  {n:4} {k}")


if __name__ == "__main__":
    main()
